import openpyxl

# Load and debug the demo file
file_path = r"D:\FREELANCERS\PAOLAG\FinAI Agent\DEMO_FINANCIAL_STATEMENT.xlsx"

wb = openpyxl.load_workbook(file_path, data_only=True)

print("=== DEBUGGING EXCEL FILE ===\n")

# Check BG sheet
if 'BG' in wb.sheetnames:
    ws = wb['BG']
    print("BG Sheet - Balance General:")
    print(f"  Company (A5): {ws['A5'].value}")
    print(f"\n  Assets:")
    print(f"    Cash (B11): {ws['B11'].value}")
    print(f"    Receivables (B12): {ws['B12'].value}")
    print(f"    Inventory (B13): {ws['B13'].value}")
    print(f"    Current Assets (B21): {ws['B21'].value}")
    print(f"    Fixed Assets (B23): {ws['B23'].value}")
    print(f"    Total Assets (B31): {ws['B31'].value}")
    print(f"\n  Liabilities:")
    print(f"    Payables (B33): {ws['B33'].value}")
    print(f"    Current Liab (B42): {ws['B42'].value}")
    print(f"    LT Liab (B46): {ws['B46'].value}")
    print(f"    Total Liab (B54): {ws['B54'].value}")
    print(f"\n  Equity:")
    print(f"    Total Equity (B62): {ws['B62'].value}")
    
    print("\n  Latest year (E column - 2024):")
    print(f"    Total Assets (E31): {ws['E31'].value}")
    print(f"    Current Assets (E21): {ws['E21'].value}")

# Check ER sheet
if 'ER' in wb.sheetnames:
    ws = wb['ER']
    print("\n\nER Sheet - Income Statement:")
    print(f"  Revenue (B10): {ws['B10'].value}")
    print(f"  COGS (B12): {ws['B12'].value}")
    print(f"  Gross Profit (B14): {ws['B14'].value}")
    print(f"  Operating Exp (B16): {ws['B16'].value}")
    print(f"  EBITDA (B21): {ws['B21'].value}")
    print(f"  Interest Exp (B27): {ws['B27'].value}")
    print(f"  Net Profit (B36): {ws['B36'].value}")
    
    print("\n  Latest year (E column - 2024):")
    print(f"    Revenue (E10): {ws['E10'].value}")
    print(f"    Net Profit (E36): {ws['E36'].value}")

print("\n=== END DEBUG ===")
