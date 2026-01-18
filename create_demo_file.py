import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Create a new workbook
wb = openpyxl.Workbook()

# Remove default sheet
wb.remove(wb.active)

# Create BG (Balance General) sheet
bg = wb.create_sheet("BG")
bg['A5'] = "Prospecto: Demo Company SA de CV"
bg['A5'].font = Font(bold=True, size=14)

# Headers for years
years = [2021, 2022, 2023, 2024]
for idx, year in enumerate(years):
    col = chr(66 + idx)  # B, C, D, E
    bg[f'{col}8'] = year
    bg[f'{col}8'].font = Font(bold=True)
    bg[f'{col}8'].alignment = Alignment(horizontal='center')

# Balance Sheet Data
# Assets
bg['A10'] = "ACTIVO"
bg['A10'].font = Font(bold=True)
bg['A11'] = "Efectivo"
bg['A12'] = "Clientes"
bg['A13'] = "Inventario"
bg['A21'] = "Total activo circulante"
bg['A21'].font = Font(bold=True)
bg['A23'] = "Activo Fijo"
bg['A31'] = "Total del activo"
bg['A31'].font = Font(bold=True, size=12)

# Sample data for assets (growing company)
assets_data = {
    11: [500000, 650000, 800000, 950000],   # Efectivo
    12: [1200000, 1400000, 1600000, 1800000],  # Clientes
    13: [800000, 900000, 1000000, 1100000],    # Inventario
    21: [2500000, 2950000, 3400000, 3850000],  # Total circulante
    23: [3500000, 3800000, 4100000, 4400000],  # Activo fijo
    31: [6000000, 6750000, 7500000, 8250000],  # Total activo
}

for row, values in assets_data.items():
    for idx, value in enumerate(values):
        col = chr(66 + idx)
        bg[f'{col}{row}'] = value
        bg[f'{col}{row}'].number_format = '#,##0'

# Liabilities
bg['A32'] = "PASIVO"
bg['A32'].font = Font(bold=True)
bg['A33'] = "Proveedores"
bg['A42'] = "Total pasivo circulante"
bg['A42'].font = Font(bold=True)
bg['A46'] = "Total pasivo largo plazo"
bg['A46'].font = Font(bold=True)
bg['A54'] = "Total del pasivo"
bg['A54'].font = Font(bold=True, size=12)

liabilities_data = {
    33: [600000, 650000, 700000, 750000],   # Proveedores
    42: [1200000, 1300000, 1400000, 1500000],  # Pasivo circulante
    46: [1800000, 1850000, 1900000, 1950000],  # Pasivo LP
    54: [3000000, 3150000, 3300000, 3450000],  # Total pasivo
}

for row, values in liabilities_data.items():
    for idx, value in enumerate(values):
        col = chr(66 + idx)
        bg[f'{col}{row}'] = value
        bg[f'{col}{row}'].number_format = '#,##0'

# Equity
bg['A55'] = "CAPITAL CONTABLE"
bg['A55'].font = Font(bold=True)
bg['A62'] = "Total capital contable"
bg['A62'].font = Font(bold=True, size=12)

equity_data = {
    62: [3000000, 3600000, 4200000, 4800000],  # Total equity
}

for row, values in equity_data.items():
    for idx, value in enumerate(values):
        col = chr(66 + idx)
        bg[f'{col}{row}'] = value
        bg[f'{col}{row}'].number_format = '#,##0'

# Create ER (Estado de Resultados) sheet
er = wb.create_sheet("ER")
er['A5'] = "Estado de Resultados"
er['A5'].font = Font(bold=True, size=14)

# Headers
for idx, year in enumerate(years):
    col = chr(66 + idx)
    er[f'{col}8'] = year
    er[f'{col}8'].font = Font(bold=True)
    er[f'{col}8'].alignment = Alignment(horizontal='center')

# Income Statement labels
er['A10'] = "Ingresos"
er['A10'].font = Font(bold=True)
er['A12'] = "Costo de Ventas"
er['A14'] = "Utilidad Bruta"
er['A14'].font = Font(bold=True)
er['A16'] = "Gastos de operación"
er['A21'] = "EBITDA"
er['A21'].font = Font(bold=True)
er['A27'] = "Gastos financieros"
er['A36'] = "Utilidad Neta"
er['A36'].font = Font(bold=True, size=12)

# Income statement data (profitable, growing company)
income_data = {
    10: [8000000, 9500000, 11000000, 12500000],  # Revenue
    12: [4800000, 5700000, 6600000, 7500000],    # COGS (60%)
    14: [3200000, 3800000, 4400000, 5000000],    # Gross profit
    16: [1600000, 1900000, 2200000, 2500000],    # Operating expenses
    21: [1600000, 1900000, 2200000, 2500000],    # EBITDA
    27: [200000, 210000, 220000, 230000],        # Interest expense
    36: [1000000, 1200000, 1400000, 1600000],    # Net profit
}

for row, values in income_data.items():
    for idx, value in enumerate(values):
        col = chr(66 + idx)
        er[f'{col}{row}'] = value
        er[f'{col}{row}'].number_format = '#,##0'

# Save the file
output_path = r"D:\FREELANCERS\PAOLAG\FinAI Agent\DEMO_FINANCIAL_STATEMENT.xlsx"
wb.save(output_path)

print(f"✅ Demo Excel file created: {output_path}")
print("\n📊 Company Profile:")
print("  - Name: Demo Company SA de CV")
print("  - Revenue (2024): $12,500,000")
print("  - Net Profit (2024): $1,600,000")
print("  - Total Assets (2024): $8,250,000")
print("  - Equity (2024): $4,800,000")
print("\n💡 Expected Analysis:")
print("  - Current Ratio: ~2.57 (Excellent)")
print("  - ROE: ~33% (Excellent)")
print("  - Debt-to-Assets: ~42% (Moderate)")
print("  - Profit Margin: ~12.8% (Good)")
print("  - Expected Category: A or B")
print("  - Expected Recommendation: APPROVE")
