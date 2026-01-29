import openpyxl
from openpyxl.styles import Font
import os

def create_scenario_excel(filename, company_name, financial_profile, metadata=None):
    wb = openpyxl.Workbook()
    
    # --- BG Sheet (Balance General) ---
    ws_bg = wb.active
    ws_bg.title = "BG"
    
    # Basic Headers
    ws_bg['A1'] = "MOSKALTI CAPITAL"
    ws_bg['A1'].font = Font(bold=True, size=14)
    ws_bg['A4'] = f"Prospecto: {company_name}"
    
    # Metadata (Optional)
    if metadata:
        ws_bg['A5'] = f"Industry: {metadata.get('industry', '')}"
        ws_bg['B5'] = metadata.get('industry', '')
        ws_bg['A6'] = f"Years in Business: {metadata.get('years_in_business', '')}"
        ws_bg['B6'] = metadata.get('years_in_business', '')
        ws_bg['A7'] = f"Requested Amount: {metadata.get('requested_amount', '')}"
        ws_bg['B7'] = metadata.get('requested_amount', '')
        ws_bg['A8'] = f"Loan Term: {metadata.get('loan_term', '')}"
        ws_bg['B8'] = metadata.get('loan_term', '')
        ws_bg['A9'] = f"Interest Rate: {metadata.get('interest_rate', '')}"
        ws_bg['B9'] = metadata.get('interest_rate', '')
        ws_bg['A10'] = f"Credit Type: {metadata.get('credit_type', '')}"
        ws_bg['B10'] = metadata.get('credit_type', '')
        ws_bg['A11'] = f"Collateral Type: {metadata.get('collateral_type', '')}"
        ws_bg['B11'] = metadata.get('collateral_type', '')

    # Years
    years = [2021, 2022, 2023]
    header_row = 13
    for i, year in enumerate(years):
        col = openpyxl.utils.get_column_letter(i + 2)
        ws_bg[f'{col}{header_row}'] = year
        ws_bg[f'{col}{header_row}'].font = Font(bold=True)
    
    # Balance Sheet Rows
    rows = [
        ("EFECTIVO Y EQUIVALENTES", "cash"),
        ("CLIENTES", "accounts_receivable"),
        ("INVENTARIOS", "inventory"),
        ("TOTAL ACTIVO CIRCULANTE", "current_assets"),
        ("ACTIVO FIJO", "fixed_assets"),
        ("TOTAL ASSETS", "total_assets"),
        ("TOTAL PASIVO CIRCULANTE", "current_liabilities"),
        ("TOTAL DEL PASIVO", "total_liabilities"),
        ("TOTAL CAPITAL CONTABLE", "shareholder_equity")
    ]
    
    for idx, (label, key) in enumerate(rows):
        row_num = idx + header_row + 1
        ws_bg[f'A{row_num}'] = label
        for i, year in enumerate(years):
            col = openpyxl.utils.get_column_letter(i + 2)
            multiplier = 0.9 if financial_profile.get('is_declining') else (1 + (i * 0.1))
            val = financial_profile['bg'].get(key, 0) * multiplier
            ws_bg[f'{col}{row_num}'] = val

    # --- ER Sheet (Estado de Resultados) ---
    ws_er = wb.create_sheet("ER")
    ws_er['A5'] = f"Prospecto: {company_name}"
    for i, year in enumerate(years):
        col = openpyxl.utils.get_column_letter(i + 2)
        ws_er[f'{col}7'] = year
    
    er_rows = [
        ("VENTAS NETAS", "revenue"),
        ("COSTO DE VENTAS", "cost_of_goods_sold"),
        ("UTILIDAD BRUTA", "gross_profit"),
        ("EBITDA", "ebitda"),
        ("UTILIDAD NETA", "net_profit")
    ]
    
    for idx, (label, key) in enumerate(er_rows):
        row_num = idx + 10
        ws_er[f'A{row_num}'] = label
        for i, year in enumerate(years):
            col = openpyxl.utils.get_column_letter(i + 2)
            multiplier = 0.85 if financial_profile.get('is_declining') else (1 + (i * 0.15))
            val = financial_profile['er'].get(key, 0) * multiplier
            ws_er[f'{col}{row_num}'] = val

    os.makedirs('demo_reports', exist_ok=True)
    path = os.path.join('demo_reports', filename)
    wb.save(path)
    print(f"✅ Generated: {path}")

# --- SCENARIOS ---

# 1. REJECTED: Weak Liquidity, High Debt, Net Loss, No Metadata
weak_profile = {
    'is_declining': True,
    'bg': {
        'cash': 20000, 'accounts_receivable': 80000, 'inventory': 400000,
        'current_assets': 500000, 'fixed_assets': 500000, 'total_assets': 1000000,
        'current_liabilities': 625000, 'total_liabilities': 850000, 'shareholder_equity': 150000
    },
    'er': {
        'revenue': 1200000, 'cost_of_goods_sold': 1000000, 'gross_profit': 200000,
        'ebitda': 10000, 'net_profit': -50000 # Negative Profit
    }
}

# 2. APPROVED: Perfect Liquidity, Low Debt, High Profit, Full Metadata
prime_profile = {
    'is_declining': False,
    'bg': {
        'cash': 1000000, 'accounts_receivable': 1000000, 'inventory': 500000,
        'current_assets': 2500000, 'fixed_assets': 2500000, 'total_assets': 5000000,
        'current_liabilities': 800000, 'total_liabilities': 1100000, 'shareholder_equity': 3900000
    },
    'er': {
        'revenue': 15000000, 'cost_of_goods_sold': 6000000, 'gross_profit': 9000000,
        'ebitda': 7000000, 'net_profit': 6000000
    }
}

prime_meta = {
    'industry': 'Biotechnology', 'years_in_business': 12,
    'requested_amount': 3000000, 'loan_term': 24, 
    'interest_rate': 0.1125, 'credit_type': 'Line of Credit',
    'collateral_type': 1 # Mortgage
}

if __name__ == '__main__':
    create_scenario_excel("REJECTED_Weak_Financials.xlsx", "FAIL VENTURES SL", weak_profile, metadata=None)
    create_scenario_excel("APPROVED_Prime_Credit.xlsx", "PRIME GLOBAL SA", prime_profile, metadata=prime_meta)
