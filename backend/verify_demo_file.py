import openpyxl
import os

def verify_file(filename):
    path = os.path.join('demo_reports', filename)
    if not os.path.exists(path):
        print(f"File not found: {path}")
        return

    print(f"--- Verifying {filename} ---")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
        
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            print(f"\nSheet [{sheet}] First 30 rows (Col A | Col B):")
            for i in range(1, 31):
                a_val = ws[f'A{i}'].value
                b_val = ws[f'B{i}'].value
                print(f"Row {i}: {a_val} | {b_val}")
                
    except Exception as e:
        print(f"Error reading file: {e}")

if __name__ == "__main__":
    verify_file("Case_1_GradeA_Collateral1.xlsx")
